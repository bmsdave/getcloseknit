"""GLA London Schools Atlas primary-to-secondary flows (migration 058, etl_school_feeders.py)."""
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

import etl_school_feeders as feeders  # noqa: E402

try:  # the ETL reads .xlsx with openpyxl, installed in the deploy venv (not api/.venv)
    import openpyxl
except ImportError:
    openpyxl = None

NOTE = ("The information is accurate as of the {year} Spring School Census which has a reference period of "
        "January {year}.")


def workbook(path, year, rows, notes_year=None):
    book = openpyxl.Workbook()
    notes = book.active
    notes.title = "Notes"
    notes.append([])
    notes.append([None, "London Schools Atlas - Underlying Data"])
    notes.append([None, NOTE.format(year=notes_year or year)])
    data = book.create_sheet(f"Feeder_Pri2Sec_{year}_LDS")
    data.append(list(feeders.HEADER))
    for row in rows:
        data.append(list(row))
    book.save(path)


ROWS = [
    (101916, "Vicar's Green Primary School", 29, "Pri2Sec_Flow1", 138610, "Alperton Community School", 8),
    (101916, "Vicar's Green Primary School", 29, "Pri2Sec_Flow2", 101941, "The Ellen Wilkinson School for Girls", 6),
    (101916, "Vicar's Green Primary School", 29, "Pri2Sec_Flow3", 139725, "William Perkin CofE High School", 3),
    (100001, "Closed Primary", 20, "Pri2Sec_Flow1", 999999, "A School Outside London", 5),
]
CURRENT = {"101916": "101916", "138610": "138610", "101941": "101941", "139725": "139725"}


@unittest.skipIf(openpyxl is None, "openpyxl is not installed")
class ReadTests(unittest.TestCase):
    def test_reads_the_data_sheet_after_checking_the_notes_year(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "f.xlsx")
            workbook(path, 2016, ROWS)
            rows = feeders.read_workbook(path, 2016)
            self.assertEqual(rows[0], ROWS[0])
            self.assertEqual(len(rows), 4)
            with self.assertRaises(SystemExit):
                feeders.read_workbook(path, 2015)  # the notes say 2016

    def test_a_sheet_with_other_columns_is_refused(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "f.xlsx")
            workbook(path, 2016, [])
            book = openpyxl.load_workbook(path)
            book.worksheets[1]["G1"] = "Pupils"
            book.save(path)
            with self.assertRaises(SystemExit):
                feeders.read_workbook(path, 2016)


class RecordTests(unittest.TestCase):
    def test_published_flows_only_as_printed_with_current_schools(self):
        out, counts = feeders.records(ROWS, 2016, CURRENT.get)
        self.assertEqual(counts, {"primaries": 1, "unmatched": 1, "ambiguous": 0, "flows": 3,
                                  "secondaries_matched": 3})
        self.assertEqual([(r["urn"], r["flow_rank"], r["secondary_urn"], r["pupils"], r["primary_total"])
                          for r in out],
                         [("101916", 1, "138610", 8, 29), ("101916", 2, "101941", 6, 29),
                          ("101916", 3, "139725", 3, 29)])
        # 12 of 29 pupils are in flows under 3: nothing is added for them
        self.assertEqual(sum(r["pupils"] for r in out), 17)

    def test_outside_london_secondary_keeps_its_printed_name_and_no_current_school(self):
        out, _ = feeders.records(ROWS, 2016, {**CURRENT, "100001": "100002"}.get)
        row = next(r for r in out if r["reported_urn"] == "100001")
        self.assertEqual((row["urn"], row["secondary_name"], row["secondary_current_urn"]),
                         ("100002", "A School Outside London", None))

    def test_the_schools_own_row_wins_and_merged_predecessors_are_left_out(self):
        rows = [(100001, "Old", 20, "Pri2Sec_Flow1", 138610, "A", 9), (101916, "Own", 29, "Pri2Sec_Flow1", 138610, "A", 8)]
        out, counts = feeders.records(rows, 2016, {"100001": "101916", "101916": "101916", "138610": "138610"}.get)
        self.assertEqual([(r["urn"], r["reported_urn"], r["pupils"]) for r in out], [("101916", "101916", 8)])
        rows = [(100001, "Old", 20, "Pri2Sec_Flow1", 138610, "A", 9), (100002, "Old", 20, "Pri2Sec_Flow1", 138610, "A", 9)]
        out, counts = feeders.records(rows, 2016, {"100001": "101916", "100002": "101916"}.get)
        self.assertEqual((out, counts["ambiguous"]), ([], 1))

    def test_suppressed_or_inconsistent_counts_are_refused_not_filled(self):
        for bad in [(101916, "x", 29, "Pri2Sec_Flow1", 138610, "A", 2),      # under 3 is never printed
                    (101916, "x", 29, "Pri2Sec_Flow1", 138610, "A", None),
                    (101916, "x", 5, "Pri2Sec_Flow1", 138610, "A", 6),       # more than the total
                    (101916, "x", 29, "Flow1", 138610, "A", 6)]:
            with self.assertRaises(SystemExit):
                feeders.records([bad], 2016, CURRENT.get)
        with self.assertRaises(SystemExit):
            feeders.records([ROWS[0], (101916, "x", 30, "Pri2Sec_Flow2", 101941, "B", 3)], 2016, CURRENT.get)


class MigrationTests(unittest.TestCase):
    def test_table_in_migration_058_and_staged(self):
        migration = (ROOT / "migrations" / "058-school-feeder.sql").read_text()
        self.assertIn("CREATE TABLE IF NOT EXISTS school_feeder_flow (", migration)
        self.assertIn("school_feeder_flow", (ROOT / "tools" / "school_history_common.py").read_text())
        self.assertEqual(feeders.LICENCE, "OGL-2.0")
        self.assertEqual(sorted(feeders.FILES), [2014, 2015, 2016])


if __name__ == "__main__":
    unittest.main()
