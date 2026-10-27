"""Readers for councils' on-time offer documents, and the cell parsers every adapter shares.

An adapter (adapters/<slug>.py) turns one document into Offer rows using these readers; the
loader then matches names to URNs, converts distances and writes the review CSV. Readers return
plain text or lists of rows (lists of cell strings) and never guess at meaning.
"""
from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass
from html.parser import HTMLParser

MILE_METRES = 1609.344
UNIT_METRES = {"miles": MILE_METRES, "mile": MILE_METRES, "mi": MILE_METRES,
               "km": 1000.0, "kilometres": 1000.0, "metres": 1.0, "meters": 1.0, "m": 1.0}


@dataclass
class Offer:
    """One printed row: a school (or one of its criteria) and the last place offered."""
    school: str                         # name as printed
    raw_value: str                      # last-offer cell as printed: '0.634 of a mile', 'All offered'
    criterion: str | None = None        # criterion the last place went under, as printed
    rule: str | None = None             # a sub-allocation printed as its own row ('Criteria 5
                                        # (distance)', 'Band 2', 'Faith places'); None: the school
    applicants: int | None = None
    offers: int | None = None
    refused: int | None = None
    entry_point: str | None = None      # None: the document's ('reception' / 'year7'); 'year3' juniors
    site: str = "main"
    all_offered: bool | None = None     # set when the document says so outside raw_value
    source_url: str | None = None       # None: the document's url
    page: int | None = None
    method: str | None = None           # None: the document's; set when a row states its own
                                        # ('walking' for one school in a straight-line table)


# --- readers ---------------------------------------------------------------------------------

def pdf_text(path: str, layout: bool = True, raw: bool = False) -> str:
    """pdftotext (poppler) output; layout keeps table columns aligned, raw keeps stream order."""
    mode = ["-raw"] if raw else (["-layout"] if layout else [])
    return subprocess.run(["pdftotext", *mode, path, "-"], check=True, capture_output=True,
                          text=True).stdout


def pdf_pages(path: str, layout: bool = True) -> list[str]:
    """Text of each page (pdftotext separates pages with form feeds)."""
    return pdf_text(path, layout=layout).split("\f")


def pdf_tables(path: str, pages: list[int] | None = None, **settings) -> list[tuple[int, list[list[str]]]]:
    """(page number from 1, table rows) for every table pdfplumber finds; cells '' when empty."""
    import pdfplumber

    found = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            if pages and number not in pages:
                continue
            for table in page.extract_tables(settings or None):
                found.append((number, [[(cell or "").strip() for cell in row] for row in table]))
    return found


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self._row, self._cell, self._depth = [], None, None, 0

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._depth += 1
            self.tables.append([])
        elif tag == "tr" and self._depth:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.tables[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._depth:
            self._depth -= 1

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def html_tables(path: str) -> list[list[list[str]]]:
    """Every <table> as rows of cell text (whitespace collapsed); colspan is not expanded."""
    with open(path, encoding="utf-8", errors="replace") as handle:
        parser = _Tables()
        parser.feed(handle.read())
    return [table for table in parser.tables if table]


def xlsx_rows(path: str, sheet: str | None = None) -> list[list[str]]:
    """Cell values of one sheet (the first by default) as strings, '' for empty."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = book[sheet] if sheet else book.worksheets[0]
    rows = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    book.close()
    return rows


def docx_tables(path: str) -> list[list[list[str]]]:
    import docx

    return [[[cell.text.strip() for cell in row.cells] for row in table.rows]
            for table in docx.Document(path).tables]


# Scanned PDFs (format 'scanned-pdf'): set OCR_HOOK to a callable(path) -> text to enable OCR.
# Until then ocr_text refuses, and the loader records the document as not read.
OCR_HOOK = None


def ocr_text(path: str) -> str:
    if OCR_HOOK is None:
        raise NotImplementedError(f"{path}: scanned document, OCR not enabled")
    return OCR_HOOK(path)


# --- cells -------------------------------------------------------------------------------------

def clean(text: str | None) -> str:
    """Collapse whitespace and unify typographic dashes and quotes."""
    text = (text or "").replace("’", "'").replace("‘", "'").replace("–", "-").replace("—", "-")
    return " ".join(text.split())


def parse_count(cell: str | None) -> int | None:
    """'1,204' -> 1204; 'N/A', '-', '' and anything else non-numeric -> None."""
    text = clean(cell).replace(",", "")
    return int(text) if text.isdigit() else None


ALL_OFFERED = re.compile(r"\ball (?:applicants? )?(?:were )?offered\b|\ball offered\b|\bundersubscribed\b|"
                         r"\ball ?pref(?:erence)?s ?met\b|\ball on[- ]time (?:offered|received an offer)\b|"  # Kingston, Merton
                         r"^all(?: applicants)?$|"  # the whole cell under '... offered' / 'Furthest distance' (Redbridge, Sutton)
                         r"\bno (?:distance )?(?:cut[- ]?off|tie[- ]?break)|"
                         r"\ball (?:on[- ]time )?applicants (?:received|were given) a place\b|"
                         r"\ball places offered\b|\bapplicants all offered\b|"
                         r"\ball (?:on[- ]time |open )?(?:applicants|applications|students)(?: who applied)? (?:were )?offered\b",
                         re.I)
ALL_REFUSED = re.compile(r"\ball refused\b|\bnone offered\b|\bno (?:places|offers) (?:were )?made\b", re.I)
# Places given on something other than distance (aptitude tests): an outcome with no distance.
NOT_BY_DISTANCE = re.compile(r"\baptitude\b", re.I)
# 'N/A (F)': a not-applicable cell with a one- or two-letter key code the document explains
# (Hillingdon: U undersubscribed, F faith); the adapter reads the key.
EMPTY = re.compile(r"^(?:n/?a|-+|none|nil|)(?:\s*\([a-z]{1,2}\))?$", re.I)
NUMBER = re.compile(r"(?<![\d.])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?|\.\d+)\s*"
                    r"(miles?|mi|of a mile|km|kilometres|metres|meters|m)?\b", re.I)


def parse_distance(raw: str | None, units: str | None) -> tuple[int | None, bool | None]:
    """(metres, all_offered) from a last-offer cell. Units printed in the cell win over the
    document's units. Raises ValueError when the cell says neither a distance nor an outcome."""
    text = clean(raw)
    if ALL_OFFERED.search(text):
        return None, True
    if ALL_REFUSED.search(text):
        return None, False
    if EMPTY.match(text) or NOT_BY_DISTANCE.search(text):
        return None, None
    found = NUMBER.findall(text)
    if len(found) != 1:
        raise ValueError(f"distance not understood: {raw!r}")
    value, unit = found[0]
    unit = (unit or units or "").lower()
    unit = "miles" if unit == "of a mile" else unit
    if unit not in UNIT_METRES:
        raise ValueError(f"no units for distance {raw!r}")
    return round(float(value.replace(",", "")) * UNIT_METRES[unit]), None
